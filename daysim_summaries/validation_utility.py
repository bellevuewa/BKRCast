import os
import shutil
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from input_configuration import project_folder
from config import suff


def read_save_rdata(filename, variablename=None):
    df = pd.read_csv(filename)
    if variablename:
        globals()[variablename] = df
    return df


def write_tables_trip_tour(outfilename, datafile, bkrcast_folder, templatefilename, outtype):
    # Load template CSV file
    templatefile = pd.read_csv(templatefilename)
    templatefile["outsheet_orig"] = templatefile["outsheet"]

    # Load the existing Excel workbook
    wb = load_workbook(os.path.join(project_folder, "daysim_summaries", "templates", f"{outfilename}.xlsm"), keep_vba=True)
    
    # Call your custom tabulation function
    for i, s in enumerate(suff, start=0):
        purp = i  # dpurp categories
        templatefile["outsheet"] = templatefile["outsheet_orig"] + "_" + s
        tabulate_summaries(datafile[datafile["dpurp"] == purp],
                           templatefile,
                           outtype,
                           wb)

    # Save the workbook
    wb.save(os.path.join(project_folder, "daysim_summaries", "output", f"{outfilename}_{bkrcast_folder.split('_')[1]}.xlsm"))
    wb.close()


def write_tables(outfilename, datafile, bkrcast_folder, templatefilename, outtype):
    # Load template CSV file
    templatefile = pd.read_csv(templatefilename)

    # load the existing Excel workbook
    output_path = os.path.join(project_folder, "daysim_summaries", "output", f"{outfilename}_{bkrcast_folder.split('_')[1]}.xlsm")
    if not os.path.exists(output_path):
        src_path = os.path.join(project_folder, "daysim_summaries", "templates", f"{outfilename}.xlsm")
        shutil.copy(src_path, output_path)
    wb = load_workbook(output_path, keep_vba=True)

    # Call your custom tabulation function
    wb = tabulate_summaries(datafile, templatefile, outtype, wb)

    # Save the workbook
    wb.save(os.path.join(project_folder, "daysim_summaries", "output", f"{outfilename}_{bkrcast_folder.split('_')[1]}.xlsm"))
    wb.close()


def output(out, outnum, wb, outsheet, celloc, xvals=[0], yvals=[0], transpose=False):
    ws = wb[outsheet] if outsheet in wb.sheetnames else wb.create_sheet(outsheet)

    # If both xvals and yvals are provided → reshape to matrix (like cast)
    if len(xvals) > 1 and len(yvals) > 1:
        base = pd.DataFrame([(x, y) for x in xvals for y in yvals], columns=["Var1", "Var2"])
        out = pd.merge(base, out, on=["Var1", "Var2"], how="left").fillna(0)
        out_pivot = out.pivot(index="Var1", columns="Var2", values="Freq").fillna(0)
        out = out_pivot.reset_index(drop=True)

    elif len(xvals) > 1 and len(yvals) <= 1:
        base = pd.DataFrame({"Var1": xvals})
        out = pd.merge(base, out, on="Var1", how="left").fillna(0)
        out = pd.DataFrame({"val": out.iloc[:, 1].astype(float)})
        if transpose:
            out = out.T

    # Convert Excel cell location (e.g. "B4") to row and column index
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    col_letter, row_start = coordinate_from_string(celloc)
    col_start = column_index_from_string(col_letter)

    # Write the DataFrame to worksheet starting from the specified cell
    for row_idx, row in enumerate(out.itertuples(index=False), start=row_start):
        for col_idx, value in enumerate(row, start=col_start):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Optional: create a named range (Excel feature)
    regname = f"out_{outnum}"
    from openpyxl.workbook.defined_name import DefinedName
    last_row = row_start + out.shape[0] - 1
    last_col = col_start + out.shape[1] - 1
    cell_range = f"{outsheet}!{celloc}:{get_column_letter(last_col)}{last_row}"
    wb.defined_names.add(DefinedName(name=regname, attr_text=cell_range))
    return wb


def tab_crosstab(df, tabrow, wb):
    dim = tabrow['dim']
    wtvar = tabrow['weights']
    var1 = tabrow['Var1']
    xvals = list(range(int(tabrow['xvalsmin']), int(tabrow['xvalsmax']) + 1))
    if 'skipxval' in tabrow and pd.notna(tabrow['skipxval']):
        skip_vals = tabrow['skipxval']
        if isinstance(skip_vals, str):
            skip_vals = eval(skip_vals)  # expect format like "[1,2,3]"
        xvals = list(set(xvals) - set(skip_vals))

    if dim == 1:
        if pd.isna(wtvar) or wtvar == "":
            out = df[var1].value_counts().reset_index()
            out.columns = ['Var1', 'values']
        else:
            out = df.groupby(var1)[wtvar].sum().reset_index()
            out.columns = ['Var1', 'values']

        out['Var1'] = out['Var1'].astype(int)
        merged = pd.DataFrame({'Var1': xvals}).merge(out, on='Var1', how='left')
        merged['values'] = merged['values'].fillna(0)
        merged.sort_values('Var1', inplace=True)

        if 'smooth' in tabrow and pd.notna(tabrow['smooth']) and tabrow['smooth'] != "":
            merged['values'] = smooth_fun(merged['values'])

        wb = output(merged, tabrow['outnum'], wb, tabrow['outsheet'], tabrow['cell_loc'], xvals)

    elif dim == 2:
        var2 = tabrow['Var2']
        yvals = list(range(int(tabrow['yvalsmin']), int(tabrow['yvalsmax']) + 1))
        if pd.isna(wtvar) or wtvar == "":
            out = df.groupby([var1, var2]).size().reset_index(name='Freq')
        else:
            out = df.groupby([var1, var2])[wtvar].sum().reset_index(name='Freq')

        out['Var1'] = out[var1].astype(int)
        out['Var2'] = out[var2].astype(int)

        grid = pd.DataFrame([(x, y) for x in xvals for y in yvals], columns=['Var1', 'Var2'])
        out = grid.merge(out, on=['Var1', 'Var2'], how='left')
        out['Freq'] = out['Freq'].fillna(0)
        out.sort_values(['Var1', 'Var2'], inplace=True)

        if 'smooth' in tabrow and pd.notna(tabrow['smooth']) and tabrow['smooth'] != "":
            for yval in yvals:
                mask = out['Var2'] == yval
                out.loc[mask, 'Freq'] = smooth_fun(out.loc[mask, 'Freq'])

        wb = output(out, tabrow['outnum'], wb, tabrow['outsheet'], tabrow['cell_loc'], xvals, yvals)
    return wb


def tab_aggregate(df, tabrow, wb):
    aggfun = str(tabrow.get("aggfun", ""))
    var1 = str(tabrow.get("Var1", ""))
    var2 = str(tabrow.get("Var2", ""))
    wtvar = str(tabrow.get("weights", ""))
    
    # Weighted or unweighted MEAN
    if aggfun == "mean":
        df_ = df[~pd.isna(df[var1])].copy(deep=True)
        if wtvar in ("", None, "NA"):
            if var2 not in ("", None, "NA", 'nan'):
                out = df.groupby(var2, dropna=False)[var1].mean().reset_index()
                out.columns = ["Var1", "values"]
            totmean = np.average(df_[var1], weights=df_[wtvar])
        else:
            if var2 not in ("", None, "NA", 'nan'):
                grouped = df_.groupby(var2, dropna=False).apply(
                    lambda x: np.average(x[var1], weights=x[wtvar])
                ).reset_index(name="values")
                grouped.columns = ["Var1", "values"]
                out = grouped
                
            totmean = np.average(df_[var1], weights=df_[wtvar])
        
        if var2 not in ("", None, "NA", 'nan'):
            yvals = list(range(int(tabrow["yvalsmin"]), int(tabrow["yvalsmax"] + 2)))
            out.loc[len(out)] = [max(yvals), totmean]
        else:
            out = pd.DataFrame({"values": [totmean]})
            yvals = [1]
        wb = output(out[['values']], tabrow["outnum"], wb, tabrow["outsheet"], tabrow["cell_loc"], yvals=yvals)
    
    # SUM aggregation (weighted or not)
    else:
        if wtvar in ("", None, "NA"):
            out = df.groupby(var2, dropna=False)[var1].sum().reset_index()
        else:
            out = df.groupby(var2, dropna=False).apply(
                lambda x: (x[var1] * x[wtvar]).sum()
            ).reset_index(name="values")
        out.columns = ["Var1", "values"]
        yvals = list(range(int(tabrow["yvalsmin"]), int(tabrow["yvalsmax"] + 1)))
        
        if str(tabrow.get("transpose", "")) not in ("", "NA", None):
            wb = output(out, tabrow["outnum"], wb, tabrow["outsheet"], tabrow["cell_loc"], yvals=yvals, transpose=True)
        else:
            wb = output(out, tabrow["outnum"], wb, tabrow["outsheet"], tabrow["cell_loc"], yvals=yvals)
    return wb


def tabulate_summaries(fulldf, tabs, datasource, wb):
    # Filter only rows that match the current data source
    tabs = tabs[tabs['data'] == datasource]

    for _, row in tabs.iterrows():
        subsetvar = row['subsetvar']
        subsetsign = row['subsetsign']
        subsetval = row['subsetval']
        tab_type = str(row['type'])

        # Filter the data if subsetvar is provided
        if pd.notna(subsetvar) and subsetvar != "":
            df = filterdt(fulldf, subsetvar, subsetsign, subsetval)
        else:
            df = fulldf

        if not df.empty:
            if tab_type == "crosstab":
                wb = tab_crosstab(df, row, wb)
            else:
                wb = tab_aggregate(df, row, wb)
    return wb



def smooth_fun(vals):
    vals = list(vals)
    for _ in range(10):
        smoothed = [vals[0]] + [
            0.25 * vals[i - 1] + 0.5 * vals[i] + 0.25 * vals[i + 1]
            for i in range(1, len(vals) - 1)
        ] + [vals[-1]]
        vals = smoothed
    return vals


def filterdt(df, var, sign, val):
    if sign == "==" or sign == "=":
        return df[df[var] == val]
    elif sign == "!=":
        return df[df[var] != val]
    elif sign == ">":
        return df[df[var] > val]
    elif sign == "<":
        return df[df[var] < val]
    elif sign == ">=":
        return df[df[var] >= val]
    elif sign == "< =":
        return df[df[var] <= val]
    else:
        raise ValueError(f"Unsupported operator: {sign}")
